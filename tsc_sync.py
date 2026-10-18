#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tsc_sync.py — обмін даними між файлом Excel «Облік_ТСЦ.xlsx» і телефоном
через ЗАШИФРОВАНІ файли-пакети (.enc).

  python tsc_sync.py export  [--xlsx ФАЙЛ] [--out ФАЙЛ] [--with-passwords]
  python tsc_sync.py import  ЗМІНИ.enc [--xlsx ФАЙЛ]

Excel-файл під час роботи програми має бути ЗАКРИТИЙ.
"""
import argparse, getpass, hashlib, json, os, shutil, sys, uuid
from datetime import datetime
from pathlib import Path

import openpyxl
from cryptography.exceptions import InvalidTag
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

# ---------- налаштування формату ----------
MAGIC = b"TSCPK1"            # 6 байт: тип і версія формату файлу
KDF_ITER = 600_000           # ітерацій PBKDF2-HMAC-SHA256 (так само в застосунку телефона)
DEFAULT_XLSX = "Облік_ТСЦ.xlsx"
HEADER_ROW, FIRST_ROW = 3, 4          # аркуші ТСЦ: заголовки в 3-му рядку, дані з 4-го
IP_COL_NAME = "ІР адреса"
# поля, які можна змінювати на телефоні (Відділ, Телефон, ТСЦ(РСЦ) рахуються з довідника користувачів)
EDITABLE = ["№", "Ім'я обладнання", "МАС адреса", "Статус", "Тип обладнання",
            "Користувач (ПІБ)", "Тип ОС", "Виробник обладнання", "Примітки",
            "Ім'я користувача"]
PASSWORD_FIELD = "Пароль"
STATUSES = ["Увімк.", "Вимк."]
LOG_NAME = "tsc_sync_log.json"        # поруч з Excel-файлом: журнал застосованих змін


# ---------- шифрування ----------
def _key(password: str, salt: bytes) -> bytes:
    return hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, KDF_ITER, 32)

def encrypt_obj(obj, password: str) -> bytes:
    salt, iv = os.urandom(16), os.urandom(12)
    data = json.dumps(obj, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return MAGIC + salt + iv + AESGCM(_key(password, salt)).encrypt(iv, data, MAGIC)

def decrypt_obj(blob: bytes, password: str):
    if len(blob) < 6 + 16 + 12 + 16 or blob[:6] != MAGIC:
        raise ValueError("Це не пакет tsc_sync або файл пошкоджено.")
    salt, iv, ct = blob[6:22], blob[22:34], blob[34:]
    try:
        data = AESGCM(_key(password, salt)).decrypt(iv, ct, MAGIC)
    except InvalidTag:
        raise ValueError("Невірний пароль або файл змінено/пошкоджено.")
    return json.loads(data.decode("utf-8"))


# ---------- допоміжне ----------
def norm(v) -> str:
    """Значення комірки -> текст для порівняння."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()

def row_hash(values: dict, fields) -> str:
    payload = json.dumps([norm(values.get(f)) for f in fields], ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]

def to_cell(field: str, text: str):
    text = text.strip()
    if text == "":
        return None
    if field == "№" and text.isdigit():
        return int(text)
    return text

def main_sheets(wb):
    """Аркуші ТСЦ — усі, що стоять перед «Довідник ТСЦ»."""
    return wb.sheetnames[: wb.sheetnames.index("Довідник ТСЦ")]

def headers(ws, row):
    return {norm(c.value): c.column for c in ws[row] if norm(c.value)}

def read_table(ws, header_row, wanted):
    """Рядки довідника як список словників (пусті пропускаються)."""
    h = headers(ws, header_row)
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        item = {k: norm(ws.cell(r, h[k]).value) for k in wanted if k in h}
        if any(item.values()):
            out.append(item)
    return out

def read_list(ws, header_row, col):
    return [x[col] for x in read_table(ws, header_row, [col]) if x.get(col)]

def fail(msg):
    print("\nПОМИЛКА:", msg)
    sys.exit(1)

def load_wb(path: Path):
    if not path.exists():
        fail(f"Файл не знайдено: {path}")
    try:
        return openpyxl.load_workbook(path)
    except PermissionError:
        fail("Не вдається відкрити файл. Закрийте його в Excel і повторіть.")

def ask_password(confirm=False):
    pw = getpass.getpass("Пароль шифрування: ")
    if not pw:
        fail("Пароль не може бути порожнім.")
    if confirm and getpass.getpass("Повторіть пароль: ") != pw:
        fail("Паролі не збігаються.")
    return pw

def state_path(xlsx: Path):
    return xlsx.with_name(LOG_NAME)

def load_state(xlsx: Path):
    p = state_path(xlsx)
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return {"last_export_id": None, "applied": []}

def save_state(xlsx: Path, st):
    state_path(xlsx).write_text(json.dumps(st, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------- ЕКСПОРТ ----------
def cmd_export(args):
    xlsx = Path(args.xlsx)
    wb = load_wb(xlsx)
    fields = EDITABLE + ([PASSWORD_FIELD] if args.with_passwords else [])

    sheets = {}
    for name in main_sheets(wb):
        ws = wb[name]
        h = headers(ws, HEADER_ROW)
        missing = [f for f in [IP_COL_NAME] + fields if f not in h]
        if missing:
            fail(f"На аркуші «{name}» немає стовпців: {', '.join(missing)}")
        rows = []
        for r in range(FIRST_ROW, ws.max_row + 1):
            ip = norm(ws.cell(r, h[IP_COL_NAME]).value)
            if not ip:
                continue
            vals = {f: norm(ws.cell(r, h[f]).value) for f in fields}
            rows.append({"ip": ip, "v": [vals[f] for f in fields], "h": row_hash(vals, fields)})
        sheets[name] = rows

    wsr = wb["Довідник ТСЦ"]
    tsc_info = {x["Назва ТСЦ"]: {"address": x.get("Повна адреса", ""), "subnet": x.get("Підмережа", ""),
                                 "gateway": x.get("Шлюз", "")}
                for x in read_table(wsr, 2, ["Назва ТСЦ", "Повна адреса", "Підмережа", "Шлюз"])}
    pack = {
        "format": "tsc-pack", "version": 1,
        "export_id": str(uuid.uuid4()),
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "fields": fields,
        "statuses": STATUSES,
        "sheets": sheets,
        "tsc_info": tsc_info,
        "lists": {
            "equipment": read_list(wb["Довідник тип обладнання"], 2, "Тип обладнання"),
            "os": read_list(wb["Довідник тип ОС"], 2, "Тип ОС"),
            "manufacturers": read_list(wb["Довідник виробників"], 3, "Виробник обладнання"),
            "departments": read_list(wb["Довідник відділ"], 2, "Відділ"),
            "tsc": list(tsc_info.keys()),
        },
        "users": read_table(wb["Довідник користувачі"], 2, ["ПІБ", "ТСЦ", "Відділ", "Телефон", "Посада"]),
    }
    pw = ask_password(confirm=True)
    out = Path(args.out)
    out.write_bytes(encrypt_obj(pack, pw))
    st = load_state(xlsx)
    st["last_export_id"] = pack["export_id"]
    save_state(xlsx, st)
    n = sum(len(v) for v in sheets.values())
    print(f"\nГотово: {out}  (аркушів: {len(sheets)}, записів-слотів: {n}, "
          f"пароль обладнання {'ВКЛЮЧЕНО' if args.with_passwords else 'не передається'}).")
    print("Покладіть цей файл у хмару й відкрийте в застосунку на телефоні.")


# ---------- ІМПОРТ ----------
def cmd_import(args):
    xlsx, src = Path(args.xlsx), Path(args.changes)
    if not src.exists():
        fail(f"Файл змін не знайдено: {src}")
    try:
        pack = decrypt_obj(src.read_bytes(), ask_password())
    except ValueError as e:
        fail(str(e))
    if pack.get("format") != "tsc-changes" or pack.get("version") != 1:
        fail("Це не файл змін із телефона (очікується формат tsc-changes, версія 1).")

    st = load_state(xlsx)
    if pack.get("based_on_export_id") != st.get("last_export_id"):
        print("УВАГА: зміни зроблено на основі НЕ останнього експорту. "
              "Конфлікти буде перевірено, але переконайтеся, що це потрібний файл.")
    wb = load_wb(xlsx)
    fields = pack["fields"]

    # індекс: аркуш -> ІР -> номер рядка
    idx = {}
    for name in main_sheets(wb):
        ws = wb[name]; h = headers(ws, HEADER_ROW)
        idx[name] = (ws, h, {norm(ws.cell(r, h[IP_COL_NAME]).value): r
                             for r in range(FIRST_ROW, ws.max_row + 1)
                             if norm(ws.cell(r, h[IP_COL_NAME]).value)})

    applied, same, skipped, conflicts, errors = [], 0, 0, [], []
    done = set(st["applied"])
    for ch in pack["changes"]:
        cid = ch["change_id"]
        if cid in done:
            skipped += 1
            continue
        if ch["sheet"] not in idx or ch["ip"] not in idx[ch["sheet"]][2]:
            errors.append(f"{ch['sheet']} / {ch['ip']}: запис не знайдено в Excel")
            continue
        ws, h, rows = idx[ch["sheet"]]
        r = rows[ch["ip"]]
        cur = {f: norm(ws.cell(r, h[f]).value) for f in fields}
        new = dict(cur); new.update({k: norm(v) for k, v in ch["values"].items()})
        ok_hashes = {ch["base_hash"], *ch.get("also_ok", [])}
        if row_hash(cur, fields) in ok_hashes:
            pass                                    # на ПК рядок не змінювався — застосовуємо
        elif all(cur[f] == new[f] for f in fields):
            same += 1; applied.append(cid)          # у Excel уже саме такі значення
            continue
        else:
            conflicts.append((ch, cur, new))        # змінено і на ПК, і на телефоні
            continue
        for f, v in ch["values"].items():
            ws.cell(r, h[f]).value = to_cell(f, v)
        applied.append(cid)

    real = len(applied) - same
    if real:
        backup_dir = xlsx.parent / "backup"; backup_dir.mkdir(exist_ok=True)
        bk = backup_dir / f"{xlsx.stem}_{datetime.now():%Y%m%d_%H%M%S}{xlsx.suffix}"
        shutil.copy2(xlsx, bk)
        try:
            tmp = xlsx.with_name(xlsx.stem + "._tmp" + xlsx.suffix)
            wb.save(tmp)
            os.replace(tmp, xlsx)
        except PermissionError:
            fail("Не вдалося записати файл. Закрийте його в Excel і повторіть. Зміни НЕ застосовано.")
        print(f"Резервну копію збережено: {bk}")
    st["applied"] = sorted(done | set(applied))
    save_state(xlsx, st)

    print(f"\nЗастосовано змін: {real} | вже збігалися: {same} | пропущено (вже було): {skipped} "
          f"| конфліктів: {len(conflicts)} | помилок: {len(errors)}")
    if conflicts or errors:
        rep = xlsx.parent / f"Конфлікти_{datetime.now():%Y%m%d_%H%M%S}.txt"
        lines = []
        for ch, cur, new in conflicts:
            lines.append(f"КОНФЛІКТ  {ch['sheet']} / {ch['ip']}  (НЕ застосовано)")
            for f in fields:
                if cur[f] != new[f]:
                    lines.append(f"   {f}: у Excel зараз = «{cur[f]}», з телефона = «{new[f]}»")
        lines += [f"ПОМИЛКА   {e}" for e in errors]
        rep.write_text("\n".join(lines), encoding="utf-8")
        print(f"Звіт про конфлікти й помилки: {rep}")


def main():
    base = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    e = sub.add_parser("export")
    e.add_argument("--xlsx", default=str(base / DEFAULT_XLSX))
    e.add_argument("--out", default=str(base / "tsc_для_телефона.enc"))
    e.add_argument("--with-passwords", action="store_true", help="передати на телефон також поле «Пароль»")
    i = sub.add_parser("import")
    i.add_argument("changes")
    i.add_argument("--xlsx", default=str(base / DEFAULT_XLSX))
    args = ap.parse_args()
    {"export": cmd_export, "import": cmd_import}[args.cmd](args)


if __name__ == "__main__":
    main()
